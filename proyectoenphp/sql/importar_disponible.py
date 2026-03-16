#!/usr/bin/env python3
"""
Importa el Excel "DISPONIBLE QUE ME PASA ALVARO A DANI.xlsx" a tbl_disponible (MySQL).

Uso:
    python importar_disponible.py

Requisitos:
    pip install openpyxl mysql-connector-python

Notas:
    - Hace TRUNCATE de tbl_disponible antes de insertar (carga completa, no incremental).
    - Salta filas sin nombre_floriday (col L) NI descripcion (col O).
    - Descuento y porcentaje_ocupacion se convierten de decimal a porcentaje (0.1 -> 10.0).
    - Campo 'disponible' convierte 'SI' -> 1, otro -> 0.
    - Campos campanya_precios_espec, producto_precio_espec: 'X' -> 1, else 0.
"""

import os
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instala con: pip install openpyxl")

try:
    import mysql.connector
except ImportError:
    sys.exit("Falta mysql-connector-python. Instala con: pip install mysql-connector-python")


# ── Configuracion BD ──────────────────────────────────────────────
DB_HOST = os.getenv("DB_HOST", "127.0.0.1")
DB_PORT = int(os.getenv("DB_PORT", "3308"))
DB_USER = os.getenv("DB_USER", "root")
DB_PASS = os.getenv("DB_PASS", "")
DB_NAME = os.getenv("DB_NAME", "licitaciones")

# ── Ruta del Excel ────────────────────────────────────────────────
EXCEL_PATH = Path(__file__).parent / "DISPONIBLE QUE ME PASA ALVARO A DANI.xlsx"

# ── Mapeo Excel col (1-based) -> columna BD ───────────────────────
# Columnas que se saltan: I(9) id_articulo_agricultor, J(10) passaporte_fito,
# W(23) vacia, Z(26)/AB(28)/AD(30) diplad variantes sin columna BD,
# BC-BF(55-58) campos calculados de ocupacion.
COL_MAP = [
    # (excel_col_1based, db_column, transform)
    (1,  "foto",                       None),
    (2,  "foto1",                      None),
    (3,  "campanya_precios_espec",     "flag_x"),    # 'X' -> 1
    (4,  "producto_precio_espec",      "flag_x"),
    (5,  "codigo",                     None),
    (6,  "codigo_rach",                "to_str"),     # puede ser numerico
    (7,  "descripcion_rach",           None),
    (8,  "ean",                        "to_str"),
    # 9, 10 skip
    (11, "floricode",                  "to_str"),
    (12, "nombre_floriday",            None),
    (13, "clasificacion",              None),
    (14, "calidad",                    None),
    (15, "descripcion",               None),
    (16, "precio_coste_productor",     "decimal"),
    (17, "descuento_productor",        "pct"),        # 0.1 -> 10.00
    (18, "precio_coste_final",         "decimal"),
    (19, "tarifa_mayorista",           "decimal"),
    (20, "precio_x_unid",             "decimal"),
    (21, "precio_x_unid_diplad_m7",   "decimal"),
    (22, "precio_x_unid_almeria",     "decimal"),
    # 23 skip (empty)
    (24, "precio_t5_directo",          "decimal"),
    (25, "precio_t5_almeria",          "decimal"),
    # 26 skip (diplad t5)
    (27, "precio_t10",                 "decimal"),
    # 28 skip (diplad t10)
    (29, "precio_t15",                 "decimal"),
    # 30 skip (diplad t15)
    (31, "precio_t25",                 "decimal"),
    (32, "precio_dipladen_t25",        "decimal"),
    (33, "formato",                    None),
    (34, "tamanyo_aprox",              None),
    (35, "observaciones",              None),
    (36, "clasificacion_compra_facil", None),
    (37, "color",                      None),
    (38, "caracteristicas",            None),
    (39, "cantidades_minimas",         "int"),
    (40, "unids_x_piso",              "int"),
    (41, "unids_x_cc",                "int"),
    (42, "porcentaje_ocupacion",       "pct"),        # 1 -> 100.00
    (43, "zona",                       None),
    (44, "disponible",                 "flag_si"),    # 'SI' -> 1
    (45, "pedido_x_unid",             "int"),
    (46, "pedido_x_piso",             "int"),
    (47, "pedido_x_cc",               "int"),
    (48, "cod_productor",              "to_str"),
    (49, "cod_productor_opc2",         "to_str"),
    (50, "cod_productor_opc3",         "to_str"),
    (51, "nombre_productor",           None),
    (52, "unids_disponibles",          "int"),
    (53, "fecha_sem_produccion",       None),
    (54, "ultimo_cambio",              None),         # texto tipo 'S08-2026', se guarda tal cual
    # 55-58 skip (campos calculados)
    # 59 skip (pasado_a_freshportal eliminado)
    (60, "total_unids_x_linea",        "int"),
    (61, "incremento_precio_x_unid",   "decimal"),
]


def transform(value, kind):
    """Aplica transformacion al valor de la celda."""
    if value is None:
        return None

    if kind is None:
        v = str(value).strip() if not isinstance(value, str) else value.strip()
        return v if v else None

    if kind == "to_str":
        v = str(value).strip() if not isinstance(value, (int, float)) else str(int(value) if isinstance(value, float) and value == int(value) else value)
        return v if v else None

    if kind == "decimal":
        try:
            return round(float(value), 2)
        except (ValueError, TypeError):
            return None

    if kind == "pct":
        # 0.1 -> 10.0, 1 -> 100.0
        try:
            v = float(value)
            if v <= 1.0:
                return round(v * 100, 2)
            return round(v, 2)
        except (ValueError, TypeError):
            return None

    if kind == "int":
        try:
            return int(float(value))
        except (ValueError, TypeError):
            return None

    if kind == "flag_x":
        return 1 if str(value).strip().upper() in ("X", "1", "SI") else 0

    if kind == "flag_si":
        return 1 if str(value).strip().upper() in ("SI", "S", "1", "YES", "X") else 0

    return value


def main():
    if not EXCEL_PATH.exists():
        sys.exit(f"No se encuentra el Excel: {EXCEL_PATH}")

    print(f"Leyendo {EXCEL_PATH.name} ...")
    wb = openpyxl.load_workbook(str(EXCEL_PATH), read_only=True, data_only=True)
    ws = wb.active

    db_columns = [col for _, col, _ in COL_MAP]
    placeholders = ", ".join(["%s"] * len(db_columns))
    col_names = ", ".join(db_columns)
    insert_sql = f"INSERT INTO tbl_disponible ({col_names}) VALUES ({placeholders})"

    rows_to_insert = []
    skipped = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # col L (idx 11) = nombre_floriday, col O (idx 14) = descripcion
        nombre = row[11] if len(row) > 11 else None
        desc = row[14] if len(row) > 14 else None

        if not nombre and not desc:
            skipped += 1
            continue

        values = []
        for excel_col, _db_col, kind in COL_MAP:
            idx = excel_col - 1
            cell_val = row[idx] if idx < len(row) else None
            values.append(transform(cell_val, kind))

        rows_to_insert.append(tuple(values))

    wb.close()
    print(f"Filas a insertar: {len(rows_to_insert)} (saltadas: {skipped})")

    if not rows_to_insert:
        sys.exit("No hay datos para insertar.")

    # ── Conectar y cargar ─────────────────────────────────────────
    print(f"Conectando a MySQL {DB_HOST}:{DB_PORT}/{DB_NAME} ...")
    conn = mysql.connector.connect(
        host=DB_HOST,
        port=DB_PORT,
        user=DB_USER,
        password=DB_PASS,
        database=DB_NAME,
        charset="utf8mb4",
    )
    cursor = conn.cursor()

    try:
        cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
        cursor.execute("TRUNCATE TABLE tbl_disponible")
        cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
        print("Tabla tbl_disponible truncada.")

        batch_size = 100
        inserted = 0
        for i in range(0, len(rows_to_insert), batch_size):
            batch = rows_to_insert[i : i + batch_size]
            cursor.executemany(insert_sql, batch)
            inserted += len(batch)
            print(f"  Insertadas {inserted}/{len(rows_to_insert)} filas ...", end="\r")

        conn.commit()
        print(f"\nImportacion completada: {inserted} filas insertadas en tbl_disponible.")

    except Exception as e:
        conn.rollback()
        print(f"\nError durante la importacion: {e}")
        raise
    finally:
        cursor.close()
        conn.close()


if __name__ == "__main__":
    main()
